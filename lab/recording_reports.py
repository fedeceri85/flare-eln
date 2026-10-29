"""Export experiment recording tables without UI actions or HTML."""
from datetime import date
from io import BytesIO
from math import ceil
from pathlib import Path
from xml.sax.saxutils import escape

from django.utils import timezone

from .ui import recording_field_label


def recording_report_data(experiment, recordings):
    kind = experiment.recording_type
    fields = list(kind.fields.all()) if kind else []
    headings = ["Date", "Animals", *[recording_field_label(field) for field in fields]]
    widths = [16, 28, *[24 for _ in fields]]
    headings.extend(["Recording Paths", "Notes"])
    widths.extend([40, 44])
    rows = []
    for recording in recordings:
        values = []
        for field in fields:
            value = recording.values.get(field.key)
            if field.data_type == "date" and isinstance(value, str):
                try:
                    value = date.fromisoformat(value)
                except ValueError:
                    pass
            values.append(value)
        row = [recording.recording_date, "\n".join(mouse.pk for mouse in recording.mice.all()), *values]
        row.extend(["\n".join(recording.data_paths), recording.notes])
        rows.append(row)
    return {"headings": headings, "widths": widths, "rows": rows}


def display_value(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    return str(value)


def render_recording_excel(data, project_name, experiment_name):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Recordings"
    last = len(data["headings"])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
    sheet.cell(1, 1, "Recordings report").font = Font(size=16, bold=True)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last)
    sheet.cell(2, 1, f"{project_name} / {experiment_name}")
    sheet.row_dimensions[2].height = 36
    sheet.append(data["headings"])
    for row in data["rows"]:
        sheet.append(row)
    edge = Side(style="thin", color="8798A5")
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row >= 3:
                cell.border = Border(left=edge, right=edge, top=edge, bottom=edge)
                cell.font = Font(size=10, bold=cell.row == 3)
            if cell.row == 3:
                cell.fill = PatternFill("solid", fgColor="E8EEF3")
            if isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"
    for index, width in enumerate(data["widths"], 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in range(4, sheet.max_row + 1):
        lines = max(sum(max(1, ceil(len(line) / (width - 2))) for line in display_value(cell.value).split("\n"))
                    for cell, width in zip(sheet[row], data["widths"]))
        sheet.row_dimensions[row].height = min(409, max(30, 15 * lines + 8))
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:{get_column_letter(last)}{sheet.max_row}"
    sheet.print_title_rows = "1:3"
    sheet.print_area = f"A1:{get_column_letter(last)}{sheet.max_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3 if last > 6 else sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.oddFooter.center.text = "Page &P of &N"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def render_recording_pdf(data, project_name, experiment_name):
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    fonts = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("ReportVera", "Vera.ttf"), ("ReportVeraBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    style = ParagraphStyle("RecordingCell", fontName="ReportVera", fontSize=8, leading=11, splitLongWords=True)
    bold = ParagraphStyle("RecordingHeading", parent=style, fontName="ReportVeraBold")

    def cell(value, heading=False):
        return Paragraph(escape(display_value(value)).replace("\n", "<br/>"), bold if heading else style)

    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A3 if len(data["headings"]) > 6 else A4),
                            leftMargin=28, rightMargin=28, topMargin=48, bottomMargin=35,
                            title="Recordings report", author="FLARE-ELN")
    generated = timezone.localdate().strftime("%d/%m/%Y")

    def frame(canvas, document):
        canvas.saveState()
        canvas.setFont("ReportVeraBold", 14)
        canvas.drawString(28, document.pagesize[1] - 28, "Recordings report")
        canvas.setFont("ReportVera", 8)
        canvas.drawString(28, 19, f"FLARE-ELN | Generated {generated}")
        canvas.drawRightString(document.pagesize[0] - 28, 19, f"Page {document.page}")
        canvas.restoreState()

    story = []
    # Wide schemas continue in panels, repeating Date and Animals for orientation.
    extra_columns = list(range(2, len(data["headings"])))
    panels = [extra_columns[start:start + 6] for start in range(0, len(extra_columns), 6)] or [[]]
    for index, extra in enumerate(panels):
        if story:
            story.append(PageBreak())
        columns = [0, 1, *extra]
        story.extend([cell(f"{project_name} / {experiment_name}", True), Spacer(1, 8)])
        if len(panels) > 1:
            story.extend([cell(f"Column group {index + 1} of {len(panels)}"), Spacer(1, 8)])
        rows = [[cell(data["headings"][col], True) for col in columns]]
        rows.extend([[cell(row[col]) for col in columns] for row in data["rows"]])
        widths = [data["widths"][col] for col in columns]
        table = Table(rows, colWidths=[doc.width * width / sum(widths) for width in widths],
                      repeatRows=1, splitByRow=1, splitInRow=1)
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef3")),
            ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#8798a5")),
            ("LEFTPADDING", (0, 0), (-1, -1), 6), ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        story.append(table)
    doc.build(story, onFirstPage=frame, onLaterPages=frame)
    return output.getvalue()
