"""Bounded, value-only spreadsheet parsing. Never execute spreadsheet formulas."""
import csv
import io
import math
import posixpath
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.utils.cell import get_column_letter, range_boundaries

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MAX_COLUMNS = 100


def cell_value(value):
    if isinstance(value, (date, datetime)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    if value is None:
        return ""
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValidationError("Spreadsheet contains a non-finite number.")
        if value.is_integer():
            return int(value)
    if isinstance(value, str) and len(value) > 10000:
        raise ValidationError("A spreadsheet cell exceeds 10,000 characters.")
    return value


def bounded_rows(rows):
    result = []
    for number, row in enumerate(rows, 1):
        values = [cell_value(value) for value in row]
        while values and values[-1] == "":
            values.pop()
        if len(values) > MAX_COLUMNS or number > MAX_ROWS:
            raise ValidationError("Import at most 5,000 rows and 100 columns per sheet.")
        result.append(values)
    return result


def read_upload(upload):
    """Return JSON-safe source data; uploaded bytes are not retained."""
    raw = upload.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValidationError("Upload a file smaller than 10 MB.")
    suffix = Path(upload.name).suffix.lower()
    try:
        if suffix == ".csv":
            text = raw.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            return [{"name": "CSV", "rows": bounded_rows(csv.reader(io.StringIO(text), dialect)), "merges": []}]
        if suffix != ".xlsx":
            raise ValidationError("Choose an .xlsx or UTF-8 CSV file.")
        with ZipFile(io.BytesIO(raw)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValidationError("The expanded workbook exceeds 50 MB.")
            ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
            rels = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            targets = {r.get("Id"): r.get("Target") for r in rels}
            meta = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            merges = {}
            for sheet in meta.findall("s:sheets/s:sheet", ns):
                target = targets[sheet.get(f"{{{rel_ns}}}id")]
                path = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
                root = ElementTree.fromstring(archive.read(path))
                merges[sheet.get("name")] = [m.get("ref") for m in root.findall("s:mergeCells/s:mergeCell", ns)]
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
        try:
            if len(workbook.worksheets) > 20:
                raise ValidationError("Import a workbook with at most 20 worksheets.")
            sheets = []
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(
                    max_row=min(sheet.max_row or 1, MAX_ROWS + 1),
                    max_col=min(sheet.max_column or 1, MAX_COLUMNS + 1), values_only=True,
                )
                sheets.append({"name": sheet.title, "rows": bounded_rows(rows), "merges": merges.get(sheet.title, [])})
            return sheets
        finally:
            workbook.close()
    except (BadZipFile, KeyError, ValueError, TypeError, IndexError, UnicodeError, csv.Error, ElementTree.ParseError, OSError, InvalidFileException) as exc:
        raise ValidationError("This file could not be read. Choose a valid .xlsx or UTF-8 CSV file.") from exc


def table_from_sheet(sheet, header_row):
    rows = sheet["rows"]
    if not 1 <= header_row <= len(rows):
        raise ValidationError("Choose an existing row for the column headings.")
    width = max((len(row) for row in rows), default=0)
    header = list(rows[header_row - 1]) + [""] * width
    for merge in sheet.get("merges", []):
        left, top, right, bottom = range_boundaries(merge)
        if top == bottom == header_row:
            for index in range(left - 1, min(right, width)):
                header[index] = header[left - 1]
    seen = Counter()
    columns = []
    for index in range(width):
        heading = str(header[index]).strip()
        seen[heading] += 1
        columns.append({
            "heading": heading, "occurrence": seen[heading], "index": index,
            "label": f"{heading or 'Unnamed'} [{get_column_letter(index + 1)}]",
        })
    data = []
    for number, values in enumerate(rows[header_row:], header_row + 1):
        values = values + [""] * (width - len(values))
        populated = sum(v != "" for v in values)
        merged_note = any(
            top <= number <= bottom and right > left
            for left, top, right, bottom in (range_boundaries(m) for m in sheet.get("merges", []))
        )
        flag = "Empty row" if not populated else "Possible metadata row" if merged_note or populated == 1 else ""
        data.append({"number": number, "values": values, "flag": flag})
    return {"columns": columns, "rows": data}


def binding(column):
    """Positions belong to previews; only heading and occurrence are portable."""
    return {"heading": column["heading"], "occurrence": column["occurrence"]}


def resolve_binding(columns, definition):
    matches = [c for c in columns if binding(c) == definition]
    if len(matches) != 1:
        raise ValidationError(f"Column {definition.get('heading', '')!r} changed or is missing. Remap it before importing.")
    return matches[0]["index"]
