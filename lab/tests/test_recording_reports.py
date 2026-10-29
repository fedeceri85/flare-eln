from datetime import date, datetime
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import Client, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from ..models import Experiment, Mouse, Project, Protocol, Recording, RecordingField, RecordingType
from ..recording_reports import render_recording_pdf
from ..test_navigation import Markup
from ._support import AuthenticatedTestCase


@override_settings(ALLOWED_HOSTS=["testserver"])
class RecordingReportTests(AuthenticatedTestCase):
    def setUp(self):
        self.kind = RecordingType.objects.create(name="Export type", slug="export-type")
        for key, label, datatype in (("count", "Count", "integer"), ("passed", "Passed", "boolean"), ("check_date", "Check date", "date")):
            RecordingField.objects.create(recording_type=self.kind, key=key, label=label, data_type=datatype, display_in_table=True)
        RecordingField.objects.create(recording_type=self.kind, key="hidden", label="Hidden", data_type="text", display_in_table=False)
        RecordingField.objects.create(recording_type=self.kind, key="retired", label="Retired", data_type="text", active=False, display_in_table=False)
        self.project = Project.objects.create(name="Export project")
        self.experiment = Experiment.objects.create(project=self.project, name="Export experiment", recording_type=self.kind)
        other_experiment = Experiment.objects.create(project=self.project, name="Other experiment", recording_type=self.kind)
        other_user = get_user_model().objects.create_user(username="other-recording-owner")
        self.recordings = [Recording.objects.create(
            experiment=self.experiment, recording_type=self.kind, owner=owner,
            recording_date=date(2026, 10, day), data_paths=[f"/data/{day}"], notes=f"Notes {day}",
            values={"count": day, "passed": False, "check_date": "2026-10-01", "hidden": "Hidden value", "retired": "Historical value"},
        ) for day, owner in ((1, self.user), (2, other_user))]
        self.foreign = Recording.objects.create(experiment=other_experiment, recording_type=self.kind, owner=self.user, recording_date=date(2026, 10, 3))
        protocol = Protocol.objects.create(protocol_number=142, licence_reference="EXPORT", name="Export")
        mouse = Mouse.objects.create(mouse_id="Animal / 007", protocol=protocol)
        self.recordings[0].mice.add(mouse)
        self.url = reverse("lab:recording_report", args=[self.project.pk, self.experiment.pk])
        self.table_url = reverse("lab:recordings_page", args=[self.project.pk, self.experiment.pk])

    def download(self, **kwargs):
        return self.client.post(self.url, {"download": "1", "format": "xlsx", "scope": "all", **kwargs})

    def test_export_button_submits_checked_recordings_including_other_owners(self):
        response = self.client.get(self.table_url)
        markup = Markup(response.content.decode())
        export = markup.find("button", formaction=self.url)[0]
        self.assertEqual(export.attrs["form"], "selected-recording-form")
        checkboxes = markup.find("input", name="selected_recording")
        self.assertEqual({node.attrs["value"] for node in checkboxes}, {str(recording.pk) for recording in self.recordings})
        response = self.client.post(self.url, {"configure_export": "1", "selected_recording": [recording.pk for recording in self.recordings]})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["form"]["scope"].value(), "selected")
        self.assertEqual(response.context["selected_count"], 2)
        self.assertContains(response, "PDF")
        self.assertContains(response, "Excel (.xlsx)")

    def test_entire_experiment_ignores_filters_and_page_and_keeps_table_sort(self):
        response = self.download(return_url=self.table_url + "?date_from=2026-10-02&page=2&sort=recording_date&direction=asc")
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 5)
        headings = [cell.value for cell in sheet[3]]
        self.assertEqual(headings[:2], ["Date", "Animals"])
        self.assertIn("Hidden", headings)
        self.assertIn("Retired", headings)
        self.assertNotIn("Actions", headings)
        self.assertEqual(sheet["A4"].value, datetime(2026, 10, 1))
        self.assertEqual(sheet["A5"].value, datetime(2026, 10, 2))
        self.assertEqual(sheet["B4"].value, "Animal / 007")
        values = dict(zip(headings, [cell.value for cell in sheet[4]]))
        self.assertEqual(values["Count"], 1)
        self.assertIs(values["Passed"], False)
        self.assertEqual(values["Check date"], datetime(2026, 10, 1))
        self.assertEqual(values["Recording Paths"], "/data/1")
        self.assertEqual(values["Notes"], "Notes 1")
        self.assertEqual(values["Hidden"], "Hidden value")
        self.assertEqual(values["Retired"], "Historical value")
        self.assertEqual(sheet["A4"].number_format, "dd/mm/yyyy")
        self.assertEqual(sheet.freeze_panes, "C4")
        self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assertIn(".xlsx", response["Content-Disposition"])

    def test_selected_only_deduplicates_and_allows_export_of_another_owners_readable_recording(self):
        recording = self.recordings[1]
        response = self.download(scope="selected", selected_recording=[recording.pk, recording.pk])
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(sheet["A4"].value, datetime(2026, 10, 2))
        deletion = reverse("lab:delete_selected_recordings", args=[self.project.pk, self.experiment.pk])
        self.assertEqual(self.client.post(deletion, {"selected_recording": recording.pk, "confirm_delete": "yes"}).status_code, 403)
        self.assertTrue(Recording.objects.filter(pk=recording.pk).exists())

    def test_empty_invalid_and_foreign_selection(self):
        self.assertContains(self.download(scope="selected"), "Select at least one recording", status_code=400)
        for value in ("bad", "-1", str(self.foreign.pk), "999999999"):
            with self.subTest(value=value):
                self.assertEqual(self.download(scope="selected", selected_recording=value).status_code, 404)
        self.assertEqual(self.download(format="csv").status_code, 400)
        self.assertEqual(self.download(scope="invalid").status_code, 400)
        wrong = reverse("lab:recording_report", args=[self.project.pk + 100, self.experiment.pk])
        self.assertEqual(self.client.get(wrong).status_code, 404)

    def test_get_defaults_to_entire_experiment_and_safe_return(self):
        response = self.client.get(self.url, {"return_url": "https://outside.example/"})
        self.assertEqual(response.context["form"]["scope"].value(), "all")
        self.assertEqual(response.context["return_url"], self.table_url)
        self.assertEqual(response.context["total_count"], 2)
        self.assertEqual(response.context["selected_count"], 0)

    def test_column_visibility_and_literal_formula_text(self):
        self.kind.hide_notes_in_table = self.kind.hide_data_paths_in_table = True
        self.kind.save()
        recording = self.recordings[0]
        recording.values["count"] = 0
        recording.save()
        self.project.name = '=HYPERLINK("https://example.invalid")'
        self.project.save()
        response = self.download()
        sheet = load_workbook(BytesIO(response.content)).active
        headings = [cell.value for cell in sheet[3]]
        self.assertIn("Notes", headings)
        self.assertIn("Recording Paths", headings)
        self.assertEqual(sheet["A2"].data_type, "s")
        self.assertIn(self.project.name, sheet["A2"].value)
        self.assertEqual(sheet.cell(5, headings.index("Count") + 1).value, 0)

    def test_pdf_and_long_notes_wide_schemas(self):
        response = self.download(format="pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF-"))
        data = {"headings": ["Date", "Animals", *[f"Field {i}" for i in range(10)]],
                "widths": [16, 28, *([24] * 10)],
                "rows": [[date(2026, 10, 1), "Animal A", *(["Long <notes> & café.\n" * 100] * 10)]]}
        self.assertTrue(render_recording_pdf(data, "Project", "Experiment").startswith(b"%PDF-"))

    def test_exports_require_authentication_and_csrf(self):
        client = Client(enforce_csrf_checks=True)
        self.assertEqual(client.get(self.url).status_code, 302)
        client.force_login(self.user)
        self.assertEqual(client.post(self.url, {"download": "1", "scope": "all", "format": "pdf"}).status_code, 403)

    def test_empty_experiment_returns_a_useful_error(self):
        self.experiment.recordings.all().delete()
        self.assertContains(self.download(), "There are no recordings to export", status_code=400)
