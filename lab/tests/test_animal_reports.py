from datetime import date, datetime
from io import BytesIO

from django.test import Client, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from ..animal_reports import HEADINGS, report_groups, render_pdf
from ..models import GenotypeCategory, Mouse, MouseGenotype, Protocol, SharedLink
from ..sharing import share_animals
from ._support import AuthenticatedTestCase, create_test_crossing


@override_settings(ALLOWED_HOSTS=["testserver"])
class AnimalReportTests(AuthenticatedTestCase):
    def setUp(self):
        self.protocol = Protocol.objects.create(protocol_number=91, licence_reference="REPORT", name="Report protocol")
        self.crossing, self.lines = create_test_crossing(("Report gene",))
        self.dad = Mouse.objects.create(mouse_id="DAD-1", protocol=self.protocol)
        self.mouse = Mouse.objects.create(
            mouse_id="Animal / 001", tattoo="007", sex="F", colour="Black and white",
            date_of_birth=date(2026, 1, 2), breeding_pair="Pair A", father=self.dad,
            mother_external_id="External mum", protocol=self.protocol, crossing_definition=self.crossing,
            status="Culled", cull_date=date(2026, 10, 29), notes="Café <healthy> & calm",
        )
        MouseGenotype.objects.create(mouse=self.mouse, mouse_line=self.lines[0], zygosity="het", is_uncertain=True)
        self.other = Mouse.objects.create(mouse_id="Other strain", crossing="Legacy strain", genotype="Legacy call", protocol=self.protocol)
        self.url = reverse("lab:animal_report")

    def test_report_requires_login_and_csrf(self):
        self.client.logout()
        self.assertEqual(self.client.get(self.url, {"selected_mouse": self.mouse.pk}).status_code, 302)
        self.assertEqual(self.client.post(self.url, {"selected_mouse": self.mouse.pk, "format": "pdf"}).status_code, 302)
        client = Client(enforce_csrf_checks=True)
        client.force_login(self.user)
        self.assertEqual(client.post(self.url, {"selected_mouse": self.mouse.pk, "format": "pdf"}).status_code, 403)
        self.assertFalse(SharedLink.objects.exists())

    def test_options_preserve_selection_without_allocating_share_code(self):
        response = self.client.get(self.url, {"selected_mouse": [self.other.pk, self.mouse.pk], "return_url": "https://outside.example/"})
        self.assertContains(response, "PDF")
        self.assertContains(response, "Excel (.xlsx)")
        self.assertEqual([mouse.pk for mouse in response.context["selected_mice"]], [self.other.pk, self.mouse.pk])
        self.assertTrue(response.context["return_url"].startswith(reverse("lab:view_selected_mice_all")))
        self.assertFalse(SharedLink.objects.exists())

    def test_invalid_selection_and_format_do_not_allocate_share_codes(self):
        for payload, status in (({}, 400), ({"selected_mouse": "Missing"}, 404),
                                ({"selected_mouse": self.mouse.pk, "format": "csv"}, 400),
                                ({"selected_mouse": [f"Mouse-{i}" for i in range(501)]}, 400)):
            with self.subTest(payload=payload):
                self.assertEqual(self.client.post(self.url, payload).status_code, status)
        self.assertFalse(SharedLink.objects.exists())

    def test_excel_matches_reference_fields_and_reuses_share_code(self):
        link = share_animals([self.mouse.pk, self.other.pk], "family", self.user)
        response = self.client.post(self.url, {"selected_mouse": [self.mouse.pk, self.other.pk, self.mouse.pk], "format": "xlsx"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn('.xlsx"', response["Content-Disposition"])
        self.assertEqual(response["Cache-Control"], "private, no-store")
        workbook = load_workbook(BytesIO(response.content))
        self.assertEqual(len(workbook.worksheets), 2)
        sheet = workbook.worksheets[0]
        self.assertIn(self.crossing.display_name, sheet["A2"].value)
        self.assertIn(link.code, sheet["F2"].value)
        self.assertEqual([cell.value for cell in sheet[3]], list(HEADINGS))
        self.assertEqual(sheet["H4"].value, "Report gene")
        self.assertIn("A3:A4", str(sheet.merged_cells))
        row = [cell.value for cell in sheet[5]]
        self.assertEqual(row[:3], ["Animal / 001\nLocal ID: 007", "Female", "Black and white"])
        self.assertEqual(row[3], datetime(2026, 1, 2))
        self.assertEqual(row[4:7], ["Pair A", "DAD-1", "External mum"])
        self.assertIn("(uncertain)", row[7])
        self.assertEqual(row[8], self.mouse.notes)
        self.assertEqual(sheet["D5"].number_format, "dd/mm/yyyy")
        self.assertEqual(sheet.freeze_panes, "B5")
        self.assertEqual(workbook.worksheets[1]["H5"].value, "Legacy call")
        self.assertEqual(SharedLink.objects.count(), 1)

    def test_genotype_subcolumns_follow_line_order_with_unknown_and_uncertain_calls(self):
        crossing, lines = create_test_crossing(("Line A", "Line B", "Line C"))
        self.mouse.genotype_calls.all().delete()
        self.mouse.crossing_definition = crossing
        self.mouse.save()
        MouseGenotype.objects.create(mouse=self.mouse, mouse_line=lines[0], zygosity="het", is_uncertain=True)
        MouseGenotype.objects.create(mouse=self.mouse, mouse_line=lines[2], zygosity="hom")
        response = self.client.post(self.url, {"selected_mouse": self.mouse.pk, "format": "xlsx"})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual([sheet.cell(4, col).value for col in range(8, 11)], ["Line A", "Line B", "Line C"])
        self.assertEqual(sheet["H3"].value, "Genotype")
        self.assertIn("H3:J3", str(sheet.merged_cells))
        self.assertIn("K3:K4", str(sheet.merged_cells))
        self.assertEqual(sheet["K3"].value, "Notes")
        self.assertIn("(uncertain)", sheet["H5"].value)
        self.assertIsNone(sheet["I5"].value)
        self.assertTrue(sheet["J5"].value)
        self.assertTrue(render_pdf(report_groups([self.mouse]), "m-maple-otter").startswith(b"%PDF-"))

    def test_common_missing_result_synonyms_are_blank_but_real_results_remain(self):
        category = GenotypeCategory.objects.create(code="test-result", label="Placeholder")
        call = self.mouse.genotype_calls.get()
        call.zygosity = category.code
        call.save()
        for label in ("Unknown", "Undetermined", "None", "N/A", "n.a.", "Not done", "Not genotyped", "Not tested", "Pending", "Not available"):
            with self.subTest(label=label):
                # Unknown already exists; exercise automatic recognition by code too.
                category.label = "Placeholder" if label == "Unknown" else label
                category.save()
                call.zygosity = "unknown" if label == "Unknown" else category.code
                call.save()
                self.assertEqual(report_groups([Mouse.objects.get(pk=self.mouse.pk)])[0]["rows"][0][7], "")
        for label in ("Negative", "Null", "Wild type"):
            category.code = "test-result"
            category.label = label
            category.save()
            call.zygosity = category.code
            call.save()
            self.assertEqual(report_groups([Mouse.objects.get(pk=self.mouse.pk)])[0]["rows"][0][7], label + " (uncertain)")

    def test_excel_treats_formula_like_identifiers_and_parents_as_text(self):
        self.mouse.father = None
        self.mouse.father_external_id = '=HYPERLINK("https://example.invalid")'
        self.mouse.save()
        response = self.client.post(self.url, {"selected_mouse": self.mouse.pk, "format": "xlsx"})
        sheet = load_workbook(BytesIO(response.content)).active
        self.assertEqual(sheet["F5"].value, self.mouse.father_external_id)
        self.assertEqual(sheet["F5"].data_type, "s")

    def test_pdf_download_contains_a_real_pdf_and_selection_share_code(self):
        response = self.client.post(self.url, {"selected_mouse": self.mouse.pk, "format": "pdf"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn('.pdf"', response["Content-Disposition"])
        self.assertEqual(response["Cache-Control"], "private, no-store")
        self.assertTrue(response.content.startswith(b"%PDF-"))
        self.assertEqual(SharedLink.objects.get().animals.get().mouse_id, self.mouse.pk)

    def test_pdf_paginates_notes_longer_than_one_page(self):
        groups = report_groups([self.mouse])
        groups[0]["rows"][0] = (*groups[0]["rows"][0][:-1], "Long notes with <tags> & accents café.\n" * 150)
        self.assertTrue(render_pdf(groups, "m-maple-otter").startswith(b"%PDF-"))
