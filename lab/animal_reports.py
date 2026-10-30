"""Animal report data and standalone PDF/Excel renderers."""
from collections import OrderedDict
from io import BytesIO
from math import ceil
from pathlib import Path
import re
from xml.sax.saxutils import escape

from django.db.models import Prefetch, prefetch_related_objects
from django.utils import timezone

from .animal_summary import animal_genotype_details
from .models import CrossingLine, genotype_category_labels


HEADINGS = ("ID", "Sex", "Colour", "Date of Birth", "Breeding pair", "Dad", "Mum", "Genotype", "Notes")
WIDTHS = (28, 11, 18, 16, 20, 22, 22, 39, 44)
PDF_WIDTHS = (56, 36, 38, 60, 49, 40, 40)


def is_missing_genotype_value(value):
    token = re.sub(r"[\W_]+", "", str(value or "").casefold())
    return token in {
        "", "unknown", "undetermined", "none", "na", "notavailable", "notapplicable",
        "notdone", "notgenotyped", "ungenotyped", "nottested", "untested",
        "notdetermined", "pending", "notrecorded", "noresult", "noresults", "missing",
    }


def report_groups(mice):
    prefetch_related_objects(
        mice, "crossing_definition__species", "genotype_calls__mouse_line",
        Prefetch("crossing_definition__crossing_lines", queryset=CrossingLine.objects.select_related("line").order_by("position", "crossing_line_id")),
    )
    labels = genotype_category_labels()
    groups = OrderedDict()
    for mouse in mice:
        strain = mouse.genetics_crossing_display or "Not recorded"
        key = (mouse.crossing_definition_id, strain)
        group = groups.setdefault(key, {"strain": strain, "rows": [], "lines": OrderedDict()})
        results, legacy, _genotype = animal_genotype_details(mouse, labels)
        genotype = {}
        for result in results:
            group["lines"][result["line_id"]] = result["label"]
            genotype[result["line_id"]] = (
                "" if is_missing_genotype_value(result["code"]) or is_missing_genotype_value(result["value"])
                else result["value"] + (" (uncertain)" if result["uncertain"] else "")
            )
        if legacy and not is_missing_genotype_value(legacy):
            group["lines"][None] = "Other genotype"
            genotype[None] = legacy
        identity = mouse.pk
        if mouse.tattoo:
            identity += f"\nLocal ID: {mouse.tattoo}"
        group["rows"].append((
            identity, mouse.get_sex_display() or "Unknown", mouse.colour,
            mouse.date_of_birth, mouse.breeding_pair,
            mouse.father_id or mouse.father_external_id,
            mouse.mother_id or mouse.mother_external_id, genotype, mouse.notes,
        ))
    for group in groups.values():
        columns = group.pop("lines")
        if not columns:
            columns[None] = "Unspecified"
        group["genotype_columns"] = list(columns.values())
        group["rows"] = [(*row[:7], *(row[7].get(key, "") for key in columns), row[8]) for row in group["rows"]]
    return list(groups.values())


def render_pdf(groups, share_code):
    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

    # Embed ReportLab's bundled Unicode fonts; no host fonts or template required.
    fonts = Path(reportlab.__file__).parent / "fonts"
    for name, filename in (("ReportVera", "Vera.ttf"), ("ReportVeraBold", "VeraBd.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(fonts / filename)))
    style = ParagraphStyle("ReportCell", fontName="ReportVera", fontSize=8, leading=10, alignment=TA_CENTER, splitLongWords=True)
    heading_style = ParagraphStyle("ReportHeading", parent=style, fontName="ReportVeraBold", fontSize=8)

    def paragraph(value, heading=False):
        if hasattr(value, "strftime"):
            value = value.strftime("%d/%m/%Y")
        return Paragraph(escape(str(value or "")).replace("\n", "<br/>"), heading_style if heading else style)

    buffer = BytesIO()
    margin = 25 * mm
    doc = SimpleDocTemplate(buffer, pagesize=A4, leftMargin=margin, rightMargin=margin, topMargin=margin + 20, bottomMargin=margin + 16,
                            title="Animal summary", author="FLARE-ELN")
    generated = timezone.localdate().strftime("%d/%m/%Y")

    def page_frame(canvas, document):
        canvas.saveState()
        canvas.setFont("ReportVeraBold", 14)
        canvas.drawString(margin, document.pagesize[1] - margin, "Animal summary")
        canvas.setFont("ReportVera", 8)
        canvas.drawString(margin, margin, f"FLARE-ELN | Generated {generated}")
        canvas.drawRightString(document.pagesize[0] - margin, margin, f"Page {document.page}")
        canvas.restoreState()

    story = []
    for group in groups:
        if story:
            story.append(PageBreak())
        count = len(group["genotype_columns"])
        # Scale every column and its text together to fit the portrait page.
        natural_widths = [*PDF_WIDTHS, *([65] * count), 80]
        scale = min(1, doc.width / sum(natural_widths))
        widths = [width * scale for width in natural_widths]
        style = ParagraphStyle("ScaledReportCell", parent=style, fontSize=8 * scale, leading=10 * scale)
        heading_style = ParagraphStyle("ScaledReportHeading", parent=heading_style, fontSize=8 * scale, leading=10 * scale)
        last = 7 + count
        metadata = [""] * (last + 1)
        metadata[0] = paragraph("STRAIN: " + group["strain"], True)
        metadata[5] = paragraph("SHARE CODE: " + share_code, True)
        top = [paragraph(heading, True) for heading in HEADINGS[:7]] + [paragraph("Genotype", True)] + [""] * (count - 1) + [paragraph("Notes", True)]
        bottom = [""] * 7 + [paragraph(line, True) for line in group["genotype_columns"]] + [""]
        rows = [metadata, top, bottom]
        for row in group["rows"]:
            sex = {"Male": "M", "Female": "F", "M": "M", "F": "F"}.get(row[1], "")
            identity = row[0].replace("\nLocal ID: ", "\n", 1)
            rows.append([paragraph(value) for value in (identity, sex, *row[2:])])
        table = Table(rows, colWidths=widths, repeatRows=3, splitByRow=1, splitInRow=1)
        spans = [("SPAN", (0, 0), (4, 0)), ("SPAN", (5, 0), (last, 0))]
        spans += [("SPAN", (col, 1), (col, 2)) for col in [*range(7), last]]
        if count > 1:
            spans.append(("SPAN", (7, 1), (last - 1, 1)))
        table.setStyle(TableStyle([
            *spans,
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("VALIGN", (0, 1), (-1, 2), "MIDDLE"),
            ("BACKGROUND", (0, 1), (-1, 2), colors.HexColor("#e8eef3")),
            ("GRID", (0, 1), (-1, -1), .45, colors.HexColor("#8798a5")),
            ("LEFTPADDING", (0, 0), (-1, -1), 3 * scale), ("RIGHTPADDING", (0, 0), (-1, -1), 3 * scale),
            ("TOPPADDING", (0, 0), (-1, -1), 7 * scale), ("BOTTOMPADDING", (0, 0), (-1, -1), 7 * scale),
        ]))
        story.append(table)
    doc.build(story, onFirstPage=page_frame, onLaterPages=page_frame)
    return buffer.getvalue()


def render_excel(groups, share_code):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.page import PageMargins

    workbook = Workbook()
    workbook.remove(workbook.active)
    border = Border(*(Side(style="thin", color="8798A5") for _ in range(4)))
    for index, group in enumerate(groups, 1):
        title = re.sub(r"[\\/*?:\[\]]", "-", group["strain"]).strip("'") or "Animals"
        sheet = workbook.create_sheet(f"{index} {title}"[:31])
        count = len(group["genotype_columns"])
        last = 8 + count
        last_letter = get_column_letter(last)
        widths = [*WIDTHS[:7], *([22] * count), WIDTHS[-1]]
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last)
        sheet.cell(1, 1, "Animal summary").font = Font(size=16, bold=True)
        sheet.merge_cells("A2:E2")
        sheet.merge_cells(start_row=2, start_column=6, end_row=2, end_column=last)
        sheet.cell(2, 1, "STRAIN: " + group["strain"])
        sheet.cell(2, 6, "SHARE CODE: " + share_code)
        sheet.row_dimensions[2].height = 40
        sheet.append([*HEADINGS[:7], "Genotype", *([None] * (count - 1)), "Notes"])
        sheet.append([None] * 7 + group["genotype_columns"] + [None])
        for column in [*range(1, 8), last]:
            sheet.merge_cells(start_row=3, start_column=column, end_row=4, end_column=column)
        if count > 1:
            sheet.merge_cells(start_row=3, start_column=8, end_row=3, end_column=last - 1)
        for values in group["rows"]:
            identity = values[0].replace("\nLocal ID: ", "\n", 1)
            sex = {"Male": "M", "Female": "F", "M": "M", "F": "F"}.get(values[1], "")
            sheet.append([identity, sex, *values[2:]])
        for row in sheet:
            for cell in row:
                # Animal IDs and notes are text, including leading '=' or '+'.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if cell.row >= 3:
                    cell.border = border
                    cell.font = Font(size=10, bold=cell.row in (3, 4))
                if cell.row in (3, 4):
                    cell.fill = PatternFill("solid", fgColor="E8EEF3")
        for row in range(5, sheet.max_row + 1):
            sheet.cell(row, 4).number_format = "dd/mm/yyyy"
            lines = max(
                sum(max(1, ceil(len(line) / (width - 2))) for line in str(cell.value or "").split("\n"))
                for cell, width in zip(sheet[row], widths)
            )
            sheet.row_dimensions[row].height = min(409, max(36, 15 * lines + 8))
        for col, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
        sheet.freeze_panes = "B5"
        sheet.print_title_rows = "1:4"
        sheet.print_area = f"A1:{last_letter}{sheet.max_row}"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_margins = PageMargins(left=.3, right=.3, top=.4, bottom=.4)
        sheet.oddFooter.center.text = "Page &P of &N"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
